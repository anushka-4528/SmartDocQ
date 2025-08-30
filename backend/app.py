# backend/app.py
from __future__ import annotations

import io
import os
import uuid
import csv
import time
import logging
from typing import Optional

from flask import Flask, Blueprint, request, jsonify
from flask_cors import CORS
from dotenv import load_dotenv

# --- File parsing / OCR ---
import fitz  # PyMuPDF
import docx
import openpyxl
import pytesseract
from PIL import Image

# --- Gemini (gen + embeddings) ---
import google.generativeai as genai

# --- Vector DB (Pinecone) ---
from pinecone import Pinecone, ServerlessSpec

# --- Prompt helpers ---
from prompts import build_prompt, build_context, parse_llm_json

# =========================================================
# Config
# =========================================================
load_dotenv()

GOOGLE_API_KEY     = os.getenv("GOOGLE_API_KEY")
MODEL_NAME         = os.getenv("MODEL_NAME", "gemini-1.5-flash")
TEMPERATURE        = float(os.getenv("TEMPERATURE", "0.2"))
MAX_CONTEXT_CHARS  = int(os.getenv("MAX_CHARS", "12000"))

# Vector config
PINECONE_API_KEY   = os.getenv("PINECONE_API_KEY")
PINECONE_INDEX     = os.getenv("PINECONE_INDEX", "smartdocq")
PINECONE_CLOUD     = os.getenv("PINECONE_CLOUD", "aws")
PINECONE_REGION    = os.getenv("PINECONE_REGION", "us-east-1")

TOP_K              = int(os.getenv("TOP_K", "5"))
CHUNK_SIZE         = int(os.getenv("CHUNK_SIZE", "1200"))
CHUNK_OVERLAP      = int(os.getenv("CHUNK_OVERLAP", "200"))

# Upload policy
ALLOWED_EXTS       = {".pdf", ".txt", ".docx", ".csv", ".xlsx", ".png", ".jpg", ".jpeg"}
MAX_UPLOAD_MB      = int(os.getenv("MAX_UPLOAD_MB", "20"))

# CORS
FRONTEND_ORIGIN    = os.getenv("FRONTEND_ORIGIN", "*")

# Gemini setup
genai.configure(api_key=GOOGLE_API_KEY)
EMBED_MODEL = "text-embedding-004"  # 768 dims

# App + API namespace
app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": FRONTEND_ORIGIN}})
api = Blueprint("api", __name__, url_prefix="/api/v1")

# Logging
logging.basicConfig(level=logging.INFO)
log = logging.getLogger("smartdocq")

UPLOAD_FOLDER = os.path.join(os.getcwd(), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# In-memory stores (OK for milestone/demo)
DOC_TEXTS: dict[str, str] = {}         # doc_id -> extracted text
DOC_META: dict[str, dict] = {}         # doc_id -> metadata
SESSION_MEMORY: dict[str, list[dict]] = {}  # session_id -> [{q, a}]
FEEDBACKS: list[dict] = []             # simple in-memory feedback list

# =========================================================
# Helpers: responses & validation
# =========================================================
def ok(payload: Optional[dict] = None, message: Optional[str] = None, status_code: int = 200):
    body = {"status": "ok"}
    if message:
        body["message"] = message
    if payload:
        body.update(payload)
    return jsonify(body), status_code

def err(message: str, code: Optional[str] = None, status_code: int = 400):
    body = {"status": "error", "error": message}
    if code:
        body["code"] = code
    return jsonify(body), status_code

@app.errorhandler(400)
def _h400(_e): return err("Bad request", status_code=400)

@app.errorhandler(404)
def _h404(_e): return err("Not found", status_code=404)

@app.errorhandler(500)
def _h500(_e): return err("Internal server error", status_code=500)

def _file_size(file_storage) -> int:
    """Return file size in bytes without loading whole file in memory."""
    pos = file_storage.stream.tell()
    file_storage.stream.seek(0, os.SEEK_END)
    size = file_storage.stream.tell()
    file_storage.stream.seek(pos, os.SEEK_SET)
    return size

def validate_upload(file) -> Optional[str]:
    if not file or file.filename == "":
        return "No file selected."
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_EXTS:
        return f"Unsupported type '{ext}'. Allowed: {sorted(ALLOWED_EXTS)}"
    size = _file_size(file)
    if size > MAX_UPLOAD_MB * 1024 * 1024:
        return f"File too large. Max {MAX_UPLOAD_MB} MB."
    return None

def validate_ask(data: dict) -> Optional[str]:
    q = data.get("question", "")
    if not isinstance(q, str) or not q.strip():
        return "Missing or empty 'question'."
    tk = data.get("top_k", TOP_K)
    if not isinstance(tk, int) or not (1 <= tk <= 10):
        return "'top_k' must be an integer between 1 and 10."
    return None

# =========================================================
# Vector DB helpers (Pinecone)
# =========================================================
pc = None
pindex = None
if PINECONE_API_KEY:
    try:
        pc = Pinecone(api_key=PINECONE_API_KEY)
        # Create index if missing (handy for dev)
        try:
            existing = [i["name"] for i in pc.list_indexes()]  # tolerate SDK variants
        except Exception:
            existing = [getattr(i, "name", None) for i in (pc.list_indexes() or [])]
        if PINECONE_INDEX not in existing:
            log.info("Creating Pinecone index %s", PINECONE_INDEX)
            pc.create_index(
                name=PINECONE_INDEX,
                dimension=768,           # text-embedding-004
                metric="cosine",
                spec=ServerlessSpec(cloud=PINECONE_CLOUD, region=PINECONE_REGION),
            )
        pindex = pc.Index(PINECONE_INDEX)
    except Exception as _e:
        log.exception("Pinecone init failed")
        pindex = None

def chunk_text(text: str, max_chars: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    chunks, i, n = [], 0, len(text)
    while i < n:
        j = min(i + max_chars, n)
        chunks.append(text[i:j])
        if j == n:
            break
        i = max(0, j - overlap)
    return chunks

def embed_text(text: str) -> list[float]:
    text = (text or "")[:8000]  # safety truncate
    out = genai.embed_content(model=EMBED_MODEL, content=text)
    return out["embedding"]

def embed_texts(texts: list[str]) -> list[list[float]]:
    return [embed_text(t) for t in texts]

def upsert_chunks_to_pinecone(doc_id: str, chunks: list[str], vectors: list[list[float]]):
    if not pindex:
        raise RuntimeError("Pinecone not configured. Set PINECONE_API_KEY / PINECONE_INDEX.")
    payload = []
    for i, (txt, vec) in enumerate(zip(chunks, vectors)):
        payload.append({
            "id": f"{doc_id}:{i}",
            "values": vec,
            "metadata": {"doc_id": doc_id, "seq": i, "snippet": txt[:200]},
        })
    B = 100
    for b in range(0, len(payload), B):
        pindex.upsert(vectors=payload[b:b+B])

def query_topk_from_pinecone(query_vec: list[float], top_k: int = TOP_K, filter_doc: Optional[str] = None):
    if not pindex:
        raise RuntimeError("Pinecone not configured. Set PINECONE_API_KEY / PINECONE_INDEX.")
    flt = {"doc_id": {"$eq": filter_doc}} if filter_doc else None
    res = pindex.query(vector=query_vec, top_k=top_k, include_metadata=True, filter=flt)
    matches = res.get("matches", []) or []
    return [{
        "id": m["id"],
        "score": m["score"],
        "doc_id": m["metadata"].get("doc_id"),
        "seq": m["metadata"].get("seq"),
        "snippet": m["metadata"].get("snippet", "")
    } for m in matches]

# =========================================================
# Extraction helpers
# =========================================================
def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_pdf(filepath)
    elif ext == ".txt":
        return extract_txt(filepath)
    elif ext == ".docx":
        return extract_docx(filepath)
    elif ext == ".csv":
        return extract_csv(filepath)
    elif ext == ".xlsx":
        return extract_xlsx(filepath)
    elif ext in (".jpg", ".jpeg", ".png"):
        return extract_image(filepath)
    else:
        raise ValueError("Unsupported file type")

def extract_pdf(filepath):
    text = ""
    with fitz.open(filepath) as doc:
        pages = len(doc)
        for page in doc:
            page_text = page.get_text()
            if page_text.strip():
                text += page_text
            else:
                # OCR fallback
                pix = page.get_pixmap()
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text += pytesseract.image_to_string(img)
    return text, pages

def extract_txt(filepath):
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read(), None

def extract_docx(filepath):
    d = docx.Document(filepath)
    text = "\n".join([p.text for p in d.paragraphs])
    return text, None

def extract_csv(filepath):
    lines = []
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(", ".join(row))
    return "\n".join(lines), None

def extract_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            text += ", ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text, None

def extract_image(filepath):
    img = Image.open(filepath)
    text = pytesseract.image_to_string(img)
    return text, None

# =========================================================
# Routes (under /api/v1)
# =========================================================
@api.post("/upload")
def upload_file():
    if "file" not in request.files:
        return err("No file part"), 400
    file = request.files["file"]
    v = validate_upload(file)
    if v:
        return err(v), 400

    filename = os.path.basename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)

    t0 = time.perf_counter()
    try:
        extracted, pages = extract_text(filepath)
    except Exception as e:
        log.exception("Failed to extract text")
        return err(f"Failed to extract text: {e}", status_code=500)
    t1 = time.perf_counter()

    if not (extracted or "").strip():
        return ok({"text": ""}, message="File uploaded, but no text extracted.")

    doc_id = str(uuid.uuid4())
    DOC_TEXTS[doc_id] = extracted
    size = os.path.getsize(filepath)
    DOC_META[doc_id] = {
        "filename": filename,
        "pages": pages,
        "size_bytes": size,
    }

    log.info("Extracted %s bytes in %.2fs", len(extracted), (t1 - t0))
    return ok(
        {
            "doc_id": doc_id,
            "meta": DOC_META[doc_id],
            "text": extracted,
        },
        message="File uploaded and text extracted successfully.",
    )

@api.post("/embed_doc")
def embed_doc():
    data = request.get_json() or {}
    doc_id = data.get("doc_id")
    if not doc_id or doc_id not in DOC_TEXTS:
        return err("Invalid or missing 'doc_id'."), 400
    try:
        text = DOC_TEXTS[doc_id]
        chunks = chunk_text(text, CHUNK_SIZE, CHUNK_OVERLAP)
        t0 = time.perf_counter()
        vectors = embed_texts(chunks)
        t1 = time.perf_counter()
        upsert_chunks_to_pinecone(doc_id, chunks, vectors)
        t2 = time.perf_counter()
        log.info("Embeddings: %.2fs, Upsert: %.2fs", (t1 - t0), (t2 - t1))
        return ok({"doc_id": doc_id, "chunks": len(chunks)}, message="Embeddings stored")
    except Exception as e:
        log.exception("Embedding/upsert failed")
        return err(f"Embedding/upsert failed: {e}", status_code=500)

@api.post("/retrieve_top5")
def retrieve_top5():
    data = request.get_json() or {}
    query = (data.get("query") or "").strip()
    doc_id = data.get("doc_id")
    if not query:
        return err("Missing 'query'."), 400
    try:
        q_vec = embed_text(query)
        results = query_topk_from_pinecone(q_vec, top_k=5, filter_doc=doc_id)
        return ok({"matches": results})
    except Exception as e:
        log.exception("Retrieval failed")
        return err(f"Retrieval failed: {e}", status_code=500)

@api.post("/ask")
def ask_question():
    data = request.get_json() or {}
    v = validate_ask(data)
    if v:
        return err(v), 400

    question     = data.get("question").strip()
    doc_id       = data.get("doc_id")              # optional filter
    session_id   = data.get("session_id", "default")
    style        = data.get("style", "concise")
    citation_mode= bool(data.get("citation_mode", True))
    model_name   = data.get("model", MODEL_NAME)
    temperature  = float(data.get("temperature", TEMPERATURE))
    top_k        = int(data.get("top_k", TOP_K))

    if not DOC_TEXTS:
        return err("No document uploaded yet."), 400

    try:
        # Session memory
        history = SESSION_MEMORY.setdefault(session_id, [])
        history_text = "\n".join([f"Q: {h['q']}\nA: {h['a']}" for h in history])

        # Retrieval
        candidate_chunks = None
        hits = []
        if pindex:
            q_vec = embed_text(question)
            hits = query_topk_from_pinecone(q_vec, top_k=top_k, filter_doc=doc_id)
            if hits:
                candidate_chunks = [h["snippet"] for h in hits]

        # Fallback to raw doc text if no retrieval
        if candidate_chunks:
            base_context_text = history_text
        else:
            ctx_text = DOC_TEXTS.get(doc_id) if (doc_id and doc_id in DOC_TEXTS) \
                       else (next(reversed(DOC_TEXTS.values()), ""))
            base_context_text = (ctx_text + ("\n" + history_text if history_text else "")).strip()

        # Build prompt
        context = build_context(base_context_text, candidate_chunks=candidate_chunks, max_chars=MAX_CONTEXT_CHARS)
        prompt  = build_prompt(context=context, question=question, style=style, citation_mode=citation_mode)

        # Generate
        model = genai.GenerativeModel(model_name)
        t0 = time.perf_counter()
        response = model.generate_content(prompt, generation_config={"temperature": temperature})
        t1 = time.perf_counter()
        raw_text = getattr(response, "text", "") or ""
        parsed = parse_llm_json(raw_text)

        # Save memory + create message id for feedback
        message_id = str(uuid.uuid4())
        history.append({"q": question, "a": parsed.get("answer", ""), "message_id": message_id})

        # Attach sources if we had retrieval hits
        if hits:
            parsed["sources"] = hits
        parsed["doc_id"]     = doc_id
        parsed["session_id"] = session_id
        parsed["message_id"] = message_id

        log.info("LLM gen time: %.2fs", (t1 - t0))
        return ok(parsed)
    except Exception as e:
        log.exception("Failed to get response")
        return err(f"Failed to get response: {e}", status_code=500)

@api.post("/feedback")
def feedback():
    data = request.get_json() or {}
    message_id = data.get("message_id")
    rating     = data.get("rating")
    note       = data.get("note", "")
    if not message_id or rating not in ("up", "down"):
        return err("Invalid 'message_id' or 'rating' (use 'up' or 'down')."), 400
    FEEDBACKS.append({"message_id": message_id, "rating": rating, "note": note})
    return ok({"message_id": message_id, "rating": rating, "note": note}, message="Feedback recorded")

@api.get("/docs")
def list_docs():
    docs = [{"doc_id": did, **meta} for did, meta in DOC_META.items()]
    return ok({"docs": docs})

@api.post("/reset")
def reset_state():
    DOC_TEXTS.clear()
    DOC_META.clear()
    SESSION_MEMORY.clear()
    FEEDBACKS.clear()
    return ok(message="All state cleared")

@api.get("/debug/stats")
def debug_stats():
    return ok({
        "docs": len(DOC_TEXTS),
        "meta": len(DOC_META),
        "sessions": len(SESSION_MEMORY),
        "feedbacks": len(FEEDBACKS),
        "vector_ready": bool(pindex),
        "model": MODEL_NAME,
        "index": PINECONE_INDEX,
    })

@api.get("/healthz")
def healthz():
    return ok({
        "status": "ok",
        "model": MODEL_NAME,
        "vector_ready": bool(pindex),
    })

# Register blueprint
app.register_blueprint(api)

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5001, debug=True)
